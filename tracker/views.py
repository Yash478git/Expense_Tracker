from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, redirect, render
from django.db.models import Sum
from django.db import models, transaction
import calendar
from django.contrib.auth.forms import SetPasswordForm
from django.contrib.auth.hashers import make_password

from datetime import date, timedelta
from django.utils import timezone
from django.contrib.auth.models import User
from django.contrib.auth import update_session_auth_hash

from .forms import (
    BudgetForm,
    CustomPasswordChangeForm,
    ProfileForm,
    RecurringTransactionForm,
    RegistrationForm,
    TransactionForm,
)
from .models import (
    Budget,
    Category,
    Transaction,
    Report,
    RecurringTransaction,
    UserProfile,
    EmailOTP,
)
from django.http import HttpResponse

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)

from reportlab.lib.enums import (
    TA_LEFT,
    TA_CENTER,
    TA_RIGHT,
)

from xml.sax.saxutils import escape
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter

from django.db import transaction as db_transaction

from .otp_service import send_otp, verify_otp


# =========================================================
# RECURRING TRANSACTION PROCESSOR
# =========================================================

def get_next_recurring_date(
    current_date,
    frequency,
    day_of_month
):
    """
    Calculate the next occurrence after current_date.

    For monthly schedules, the requested day is capped to
    the last valid day of the target month.

    For yearly schedules, the month is preserved and the
    year is advanced by one.
    """

    if frequency == 'yearly':
        next_year = current_date.year + 1
        target_month = current_date.month

    else:
        next_year = current_date.year
        target_month = current_date.month + 1

        if target_month > 12:
            target_month = 1
            next_year += 1

    last_day = calendar.monthrange(
        next_year,
        target_month
    )[1]

    return date(
        next_year,
        target_month,
        min(day_of_month, last_day)
    )


def process_due_recurring_transactions(user):
    """
    Materialize every due recurring transaction into a normal
    Transaction record.

    The function can safely be called repeatedly because
    next_date is advanced after every generated transaction.
    """

    today = timezone.localdate()
    created_count = 0

    recurring_queryset = (
        RecurringTransaction.objects
        .filter(
            user=user,
            is_active=True,
            next_date__lte=today
        )
        .order_by('next_date', 'id')
    )

    for recurring in recurring_queryset:

        with transaction.atomic():

            locked_recurring = (
                RecurringTransaction.objects
                .select_for_update()
                .select_related('category')
                .get(pk=recurring.pk)
            )

            # Another request may have already processed it.
            if (
                not locked_recurring.is_active
                or locked_recurring.next_date > today
            ):
                continue

            while (
                locked_recurring.is_active
                and locked_recurring.next_date <= today
            ):

                occurrence_date = (
                    locked_recurring.next_date
                )

                # Do not create an occurrence beyond the
                # configured end date.
                if (
                    locked_recurring.end_date
                    and occurrence_date
                    > locked_recurring.end_date
                ):
                    locked_recurring.is_active = False
                    locked_recurring.save(
                        update_fields=[
                            'is_active',
                            'updated_at',
                        ]
                    )
                    break

                Transaction.objects.create(
                    user=locked_recurring.user,
                    category=locked_recurring.category,
                    amount=locked_recurring.amount,
                    type=locked_recurring.type,
                    date=occurrence_date,
                    payment_method=(
                        locked_recurring.payment_method
                    ),
                    description=(
                        locked_recurring.description
                    ),
                )

                created_count += 1

                next_date = get_next_recurring_date(
                    occurrence_date,
                    locked_recurring.frequency,
                    locked_recurring.day_of_month
                )

                locked_recurring.next_date = next_date

                if (
                    locked_recurring.end_date
                    and next_date > locked_recurring.end_date
                ):
                    locked_recurring.is_active = False
                    locked_recurring.save(
                        update_fields=[
                            'next_date',
                            'is_active',
                            'updated_at',
                        ]
                    )
                    break

                locked_recurring.save(
                    update_fields=[
                        'next_date',
                        'updated_at',
                    ]
                )

    return created_count


def register(request):
    if request.method == 'POST':
        form = RegistrationForm(request.POST)

        if form.is_valid():
            email = form.cleaned_data['email']

            try:
                with db_transaction.atomic():
                    user = form.save(commit=False)

                    # Account remains inactive until email verification.
                    user.is_active = False
                    user.save()

                    UserProfile.objects.update_or_create(
                        user=user,
                        defaults={
                            'phone': form.cleaned_data['phone']
                        }
                    )

                    success, message = send_otp(
                        email=email,
                        purpose='registration',
                        user=user,
                    )

                    if not success:
                        raise ValueError(message)

                request.session['registration_user_id'] = user.id

                messages.success(
                    request,
                    'A verification OTP has been sent to your email.'
                )

                return redirect(
                    'verify_registration_otp'
                )

            except Exception:
                if 'user' in locals() and user.pk:
                    user.delete()

                messages.error(
                    request,
                    'We could not send the verification OTP. '
                    'Please try again.'
                )

    else:
        form = RegistrationForm()

    return render(
        request,
        'tracker/register.html',
        {'form': form}
    )

    return render(request, 'tracker/register.html', {'form': form})
def verify_registration_otp(request):
    user_id = request.session.get(
        'registration_user_id'
    )

    if not user_id:
        messages.error(
            request,
            'Your registration verification session has expired. '
            'Please register again.'
        )
        return redirect('register')

    user = get_object_or_404(
        User,
        pk=user_id,
        is_active=False,
    )

    if request.method == 'POST':
        otp = request.POST.get(
            'otp',
            ''
        ).strip()

        if not otp.isdigit() or len(otp) != 6:
            messages.error(
                request,
                'Please enter a valid 6-digit OTP.'
            )

            return render(
                request,
                'tracker/verify_otp.html',
                {
                    'purpose': 'registration',
                    'email': user.email,
                }
            )

        success, message, otp_record = verify_otp(
            email=user.email,
            purpose='registration',
            otp=otp,
        )

        if success:
            with db_transaction.atomic():
                user.is_active = True
                user.save(
                    update_fields=['is_active']
                )

            request.session.pop(
                'registration_user_id',
                None
            )

            messages.success(
                request,
                'Email verified successfully. '
                'Your account is now active.'
            )

            return redirect('login')

        messages.error(
            request,
            message
        )

    return render(
        request,
        'tracker/verify_otp.html',
        {
            'purpose': 'registration',
            'email': user.email,
        }
    )
def user_login(request):
    if request.method == 'POST':
        email_or_username = request.POST.get(
            'email',
            ''
        ).strip()

        password = request.POST.get(
            'password',
            ''
        )

        # Try login using username first.
        user = authenticate(
            request,
            username=email_or_username,
            password=password
        )

        # If username login fails, try matching the email.
        if user is None:
            matching_user = User.objects.filter(
                email__iexact=email_or_username
            ).first()

            if matching_user is not None:
                user = authenticate(
                    request,
                    username=matching_user.username,
                    password=password
                )

        if user is not None:
            login(request, user)

            # Administrator accounts go directly
            # to the custom admin dashboard.
            if user.is_staff:
                return redirect('admin_dashboard')

            # Regular users go to the normal dashboard.
            return redirect('dashboard')

        messages.error(
            request,
            'Invalid username/email or password.'
        )

    return render(
        request,
        'tracker/login.html'
    )

def forgot_password(request):
    if request.method == 'POST':
        email = request.POST.get(
            'email',
            ''
        ).strip().lower()

        user = User.objects.filter(
            email__iexact=email,
            is_active=True,
        ).first()

        if user is not None:
            success, message = send_otp(
                email=user.email,
                purpose='password_reset',
                user=user,
            )

            if success:
                request.session['password_reset_email'] = (
                    user.email
                )

                messages.success(
                    request,
                    'A password reset OTP has been sent to your email.'
                )

                return redirect(
                    'verify_password_reset_otp'
                )

            messages.error(
                request,
                message
            )

            return redirect('forgot_password')

        # Do not reveal whether an email address is registered.
        messages.success(
            request,
            'If an account exists for that email address, '
            'a password reset OTP has been sent.'
        )

        return redirect('forgot_password')

    return render(
        request,
        'tracker/forgot_password.html'
    )

def verify_password_reset_otp(request):
    email = request.session.get(
        'password_reset_email'
    )

    if not email:
        messages.error(
            request,
            'Your password reset session has expired. '
            'Please start again.'
        )

        return redirect('forgot_password')

    user = User.objects.filter(
        email__iexact=email,
        is_active=True,
    ).first()

    if user is None:
        messages.error(
            request,
            'Your password reset session is no longer valid.'
        )

        request.session.pop(
            'password_reset_email',
            None
        )

        return redirect('forgot_password')

    if request.method == 'POST':
        otp = request.POST.get(
            'otp',
            ''
        ).strip()

        if not otp.isdigit() or len(otp) != 6:
            messages.error(
                request,
                'Please enter a valid 6-digit OTP.'
            )

            return render(
                request,
                'tracker/verify_otp.html',
                {
                    'purpose': 'password_reset',
                    'email': user.email,
                }
            )

        success, message, otp_record = verify_otp(
            email=user.email,
            purpose='password_reset',
            otp=otp,
        )

        if success:
            request.session[
                'password_reset_verified_user_id'
            ] = user.id

            request.session.pop(
                'password_reset_email',
                None
            )

            messages.success(
                request,
                'Email verified. You can now set a new password.'
            )

            return redirect(
                'reset_password'
            )

        messages.error(
            request,
            message
        )

    return render(
        request,
        'tracker/verify_otp.html',
        {
            'purpose': 'password_reset',
            'email': user.email,
        }
    )

def reset_password(request):
    user_id = request.session.get(
        'password_reset_verified_user_id'
    )

    if not user_id:
        messages.error(
            request,
            'Your password reset session has expired. '
            'Please start again.'
        )

        return redirect('forgot_password')

    user = get_object_or_404(
        User,
        pk=user_id,
        is_active=True,
    )

    if request.method == 'POST':
        form = SetPasswordForm(
            user=user,
            data=request.POST
        )

        if form.is_valid():
            form.save()

            request.session.pop(
                'password_reset_verified_user_id',
                None
            )

            messages.success(
                request,
                'Your password has been reset successfully. '
                'You can now log in.'
            )

            return redirect('login')

    else:
        form = SetPasswordForm(
            user=user
        )

    return render(
        request,
        'tracker/reset_password.html',
        {
            'form': form,
        }
    )

@login_required
def profile(request):
    if request.method == 'POST':
        form = ProfileForm(
            request.POST,
            user=request.user
        )

        if form.is_valid():
            form.save()

            messages.success(
                request,
                'Profile updated successfully.'
            )

            return redirect('profile')

    else:
        form = ProfileForm(
            user=request.user
        )

    return render(
        request,
        'tracker/profile.html',
        {
            'form': form,
        }
    )

@login_required
def change_password(request):
    if request.method == 'POST':
        form = CustomPasswordChangeForm(
            user=request.user,
            data=request.POST
        )

        if form.is_valid():
            new_password = form.cleaned_data['new_password1']

            pending_password_hash = make_password(
                new_password
            )

            success, message = send_otp(
                email=request.user.email,
                purpose='password_change',
                user=request.user,
            )

            if not success:
                messages.error(
                    request,
                    message
                )

                return render(
                    request,
                    'tracker/change_password.html',
                    {'form': form}
                )

            otp_record = (
                EmailOTP.objects
                .filter(
                    email=request.user.email,
                    purpose='password_change',
                    is_used=False,
                )
                .order_by('-created_at')
                .first()
            )

            if otp_record is None:
                messages.error(
                    request,
                    'We could not prepare the password change verification. '
                    'Please try again.'
                )

                return render(
                    request,
                    'tracker/change_password.html',
                    {'form': form}
                )

            otp_record.pending_password_hash = (
                pending_password_hash
            )

            otp_record.save(
                update_fields=['pending_password_hash']
            )

            request.session['password_change_user_id'] = (
                request.user.id
            )

            messages.success(
                request,
                'A verification OTP has been sent to your email.'
            )

            return redirect(
                'verify_password_change_otp'
            )

    else:
        form = CustomPasswordChangeForm(
            user=request.user
        )

    return render(
        request,
        'tracker/change_password.html',
        {'form': form}
    )

@login_required
def verify_password_change_otp(request):
    user_id = request.session.get(
        'password_change_user_id'
    )

    if not user_id or user_id != request.user.id:
        messages.error(
            request,
            'Your password change verification session has expired.'
        )

        return redirect('change_password')

    user = get_object_or_404(
        User,
        pk=user_id,
        is_active=True,
    )

    if request.method == 'POST':
        otp = request.POST.get(
            'otp',
            ''
        ).strip()

        if not otp.isdigit() or len(otp) != 6:
            messages.error(
                request,
                'Please enter a valid 6-digit OTP.'
            )

            return render(
                request,
                'tracker/verify_otp.html',
                {
                    'purpose': 'password_change',
                    'email': user.email,
                }
            )

        success, message, otp_record = verify_otp(
            email=user.email,
            purpose='password_change',
            otp=otp,
        )

        if success:
            if not otp_record.pending_password_hash:
                messages.error(
                    request,
                    'The password change verification data is missing. '
                    'Please start again.'
                )

                return redirect('change_password')

            user.password = (
                otp_record.pending_password_hash
            )

            user.save(
                update_fields=['password']
            )

            update_session_auth_hash(
                request,
                user
            )

            otp_record.pending_password_hash = None
            otp_record.save(
                update_fields=['pending_password_hash']
            )

            request.session.pop(
                'password_change_user_id',
                None
            )

            messages.success(
                request,
                'Your password has been changed successfully.'
            )

            return redirect('profile')

        messages.error(
            request,
            message
        )

    return render(
        request,
        'tracker/verify_otp.html',
        {
            'purpose': 'password_change',
            'email': user.email,
        }
    )

@login_required
def delete_account(request):
    if request.user.is_staff:
        messages.error(
            request,
            'Administrator accounts cannot be deleted from this page.'
        )
        return redirect('profile')

    if request.method == 'POST':
        password = request.POST.get(
            'password',
            ''
        )

        if not request.user.check_password(password):
            messages.error(
                request,
                'Incorrect password. Account deletion was not started.'
            )

            return redirect('profile')

        success, message = send_otp(
            email=request.user.email,
            purpose='account_delete',
            user=request.user,
        )

        if not success:
            messages.error(
                request,
                message
            )

            return redirect('profile')

        request.session['delete_account_user_id'] = request.user.id

        messages.success(
            request,
            'A verification OTP has been sent to your email.'
        )

        return redirect(
            'verify_delete_account_otp'
        )

    return redirect('profile')

@login_required
def dashboard(request):
    # ---------------------------------------------------------
    # PROCESS DUE RECURRING TRANSACTIONS
    # ---------------------------------------------------------

    created_recurring_count = (
        process_due_recurring_transactions(
            request.user
        )
    )

    if created_recurring_count:
        messages.success(
            request,
            (
                f'{created_recurring_count} recurring '
                'transaction'
                f'{"s" if created_recurring_count != 1 else ""} '
                'automatically added.'
            )
        )

    # Refresh the queryset so newly created transactions
    # are included in dashboard totals and charts.
    all_transactions = request.user.transactions.all()

    # ---------------------------------------------------------
    # DASHBOARD TIME FILTER
    # ---------------------------------------------------------

    today = timezone.localdate()

    selected_period = request.GET.get(
        'period',
        'this_month'
    )

    start_date = None
    end_date = today

    if selected_period == 'this_month':
        start_date = today.replace(
            day=1
        )

    elif selected_period == 'last_month':
        if today.month == 1:
            start_date = date(
                today.year - 1,
                12,
                1
            )
        else:
            start_date = date(
                today.year,
                today.month - 1,
                1
            )

        end_date = (
            start_date + timedelta(days=32)
        ).replace(
            day=1
        ) - timedelta(days=1)

    elif selected_period == 'this_year':
        start_date = date(
            today.year,
            1,
            1
        )

    elif selected_period == 'last_year':
        start_date = date(
            today.year - 1,
            1,
            1
        )

        end_date = date(
            today.year - 1,
            12,
            31
        )

    elif selected_period == 'custom':
        custom_from = request.GET.get(
            'date_from',
            ''
        ).strip()

        custom_to = request.GET.get(
            'date_to',
            ''
        ).strip()

        try:
            start_date = date.fromisoformat(
                custom_from
            )

            end_date = date.fromisoformat(
                custom_to
            )

            if start_date > end_date:
                start_date, end_date = (
                    end_date,
                    start_date
                )

        except (ValueError, TypeError):
            start_date = today.replace(
                day=1
            )
            end_date = today

    else:
        selected_period = 'this_month'
        start_date = today.replace(
            day=1
        )

    # Transactions used by the dashboard cards/charts.
    transactions = all_transactions.filter(
        date__gte=start_date,
        date__lte=end_date
    )

    # ---------------------------------------------------------
    # BASIC FINANCIAL SUMMARY
    # ---------------------------------------------------------

    total_income = transactions.filter(
        type='income'
    ).aggregate(
        total=Sum('amount')
    )['total'] or 0

    total_expenses = transactions.filter(
        type='expense'
    ).aggregate(
        total=Sum('amount')
    )['total'] or 0

    balance = total_income - total_expenses

    # ---------------------------------------------------------
    # RECENT TRANSACTIONS
    # ---------------------------------------------------------

    recent_transactions = transactions.select_related(
        'category'
    ).order_by(
        '-date',
        '-created_at'
    )[:5]

    # ---------------------------------------------------------
    # EXPENSE BY CATEGORY
    # ---------------------------------------------------------

    expense_by_category = transactions.filter(
        type='expense'
    ).values(
        'category__name'
    ).annotate(
        total=Sum('amount')
    ).order_by(
        '-total'
    )

    # ---------------------------------------------------------
    # BUDGET ALERTS
    # ---------------------------------------------------------

    budgets = Budget.objects.filter(
        user=request.user
    ).select_related(
        'category'
    )

    exceeded_budgets = []

    for budget in budgets:
        # Budget spending must use ALL transactions so that
        # the dashboard filter does not hide budget activity.
        spent = all_transactions.filter(
            category=budget.category,
            type='expense',
            date__gte=budget.start_date,
            date__lte=budget.end_date
        ).aggregate(
            total=Sum('amount')
        )['total'] or 0

        if spent >= budget.limit:
            exceeded_budgets.append({
                'category': budget.category,
                'limit': budget.limit,
                'spent': spent,
                'remaining': budget.limit - spent,
                'period': budget.period,
            })

    # ---------------------------------------------------------
    # BUDGET PROGRESS
    # ---------------------------------------------------------

    budget_progress = []

    for budget in budgets:
        # Only show budgets that are active today.
        if not (
            budget.start_date
            <= today
            <= budget.end_date
        ):
            continue

        spent = all_transactions.filter(
            category=budget.category,
            type='expense',
            date__gte=budget.start_date,
            date__lte=budget.end_date
        ).aggregate(
            total=Sum('amount')
        )['total'] or 0

        if budget.limit > 0:
            percentage = (
                spent / budget.limit
            ) * 100
        else:
            percentage = 0

        # Keep the visual progress bar within 0–100%.
        display_percentage = min(
            round(percentage, 1),
            100
        )

        remaining = budget.limit - spent

        if percentage >= 100:
            status = 'exceeded'
        elif percentage >= 80:
            status = 'warning'
        else:
            status = 'safe'

        budget_progress.append({
            'category': budget.category,
            'limit': budget.limit,
            'spent': spent,
            'remaining': remaining,
            'percentage': display_percentage,
            'status': status,
            'period': budget.period,
            'start_date': budget.start_date,
            'end_date': budget.end_date,
        })

    # ---------------------------------------------------------
    # SMART INSIGHTS
    # ---------------------------------------------------------

    # 1. Savings Rate for the selected dashboard period.
    if total_income > 0:
        savings_rate = (
            balance / total_income
        ) * 100
    else:
        savings_rate = 0

    savings_rate = round(
        savings_rate,
        1
    )

    # 2. Top Spending Category for the selected period.
    top_spending_category = (
        expense_by_category.first()
    )

    if top_spending_category:
        top_category_name = (
            top_spending_category['category__name']
        )
        top_category_amount = (
            top_spending_category['total']
        )
    else:
        top_category_name = 'No expenses yet'
        top_category_amount = 0

    # 3. Highest Expense in the selected period.
    highest_expense = transactions.filter(
        type='expense'
    ).select_related(
        'category'
    ).order_by(
        '-amount',
        '-date'
    ).first()

    # 4. Highest Income in the selected period.
    highest_income = transactions.filter(
        type='income'
    ).select_related(
        'category'
    ).order_by(
        '-amount',
        '-date'
    ).first()

    # ---------------------------------------------------------
    # MONTH-OVER-MONTH EXPENSE COMPARISON
    # ---------------------------------------------------------

    current_month_start = date(
        today.year,
        today.month,
        1
    )

    if today.month == 1:
        previous_month_start = date(
            today.year - 1,
            12,
            1
        )
    else:
        previous_month_start = date(
            today.year,
            today.month - 1,
            1
        )

    current_month_expenses = (
        all_transactions.filter(
            type='expense',
            date__gte=current_month_start,
            date__lte=today
        ).aggregate(
            total=Sum('amount')
        )['total'] or 0
    )

    previous_month_expenses = (
        all_transactions.filter(
            type='expense',
            date__gte=previous_month_start,
            date__lt=current_month_start
        ).aggregate(
            total=Sum('amount')
        )['total'] or 0
    )

    if previous_month_expenses > 0:
        expense_change = (
            (
                current_month_expenses
                - previous_month_expenses
            )
            / previous_month_expenses
        ) * 100

        expense_change = round(
            expense_change,
            1
        )
    else:
        expense_change = 0

    # ---------------------------------------------------------
    # BUDGET UTILIZATION
    # ---------------------------------------------------------

    active_budgets = budgets.filter(
        start_date__lte=today,
        end_date__gte=today
    )

    total_budget_limit = 0
    total_budget_used = 0
    active_budget_count = 0

    for budget in active_budgets:
        spent = all_transactions.filter(
            category=budget.category,
            type='expense',
            date__gte=budget.start_date,
            date__lte=budget.end_date
        ).aggregate(
            total=Sum('amount')
        )['total'] or 0

        total_budget_limit += budget.limit
        total_budget_used += min(
            spent,
            budget.limit
        )
        active_budget_count += 1

    if total_budget_limit > 0:
        budget_utilization = (
            total_budget_used
            / total_budget_limit
        ) * 100

        budget_utilization = round(
            budget_utilization,
            1
        )
    else:
        budget_utilization = 0

    # ---------------------------------------------------------
    # SMART INSIGHTS OBJECT
    # ---------------------------------------------------------

    insights = {
        'savings_rate': savings_rate,

        'top_category_name': top_category_name,
        'top_category_amount': top_category_amount,

        'highest_expense': highest_expense,
        'highest_income': highest_income,

        'current_month_expenses': (
            current_month_expenses
        ),
        'previous_month_expenses': (
            previous_month_expenses
        ),
        'expense_change': expense_change,

        'budget_utilization': budget_utilization,
        'active_budget_count': active_budget_count,
    }

    # ---------------------------------------------------------
    # FINANCIAL HEALTH SCORE
    # ---------------------------------------------------------
    #
    # This is a transparent, rule-based application metric.
    # It is not professional financial advice.

    # Savings contribution: maximum 35 points.
    if savings_rate >= 30:
        savings_score = 35
    elif savings_rate >= 20:
        savings_score = 30
    elif savings_rate >= 10:
        savings_score = 20
    elif savings_rate > 0:
        savings_score = 10
    else:
        savings_score = 0

    # Budget utilization contribution: maximum 25 points.
    if budget_utilization <= 50:
        budget_score = 25
    elif budget_utilization <= 75:
        budget_score = 20
    elif budget_utilization <= 90:
        budget_score = 12
    elif budget_utilization <= 100:
        budget_score = 5
    else:
        budget_score = 0

    # Budget breach contribution: maximum 20 points.
    exceeded_count = len(exceeded_budgets)

    if exceeded_count == 0:
        breach_score = 20
    elif exceeded_count == 1:
        breach_score = 8
    elif exceeded_count == 2:
        breach_score = 4
    else:
        breach_score = 0

    # Expense trend contribution: maximum 20 points.
    if expense_change <= -10:
        trend_score = 20
    elif expense_change <= 0:
        trend_score = 15
    elif expense_change <= 10:
        trend_score = 10
    elif expense_change <= 20:
        trend_score = 5
    else:
        trend_score = 0

    financial_health_score = (
        savings_score
        + budget_score
        + breach_score
        + trend_score
    )

    if financial_health_score >= 85:
        financial_health_status = 'Excellent'
    elif financial_health_score >= 70:
        financial_health_status = 'Healthy'
    elif financial_health_score >= 50:
        financial_health_status = 'Moderate'
    else:
        financial_health_status = 'Needs Attention'

    financial_health = {
        'score': financial_health_score,
        'status': financial_health_status,
        'savings_score': savings_score,
        'budget_score': budget_score,
        'breach_score': breach_score,
        'trend_score': trend_score,
    }


    # ---------------------------------------------------------
    # DASHBOARD
    # ---------------------------------------------------------

    return render(
        request,
        'tracker/dashboard.html',
        {
            'total_income': total_income,
            'total_expenses': total_expenses,
            'balance': balance,
            'exceeded_budgets': exceeded_budgets,
            'expense_by_category': expense_by_category,
            'recent_transactions': recent_transactions,
            'insights': insights,
            'financial_health': financial_health,
            'budget_progress': budget_progress,

            # Time-filter values for the dashboard UI.
            'selected_period': selected_period,
            'filter_start_date': start_date,
            'filter_end_date': end_date,
        }
    )


@login_required
def admin_dashboard(request):
    if not request.user.is_staff:
        return redirect('dashboard')

    users_count = User.objects.filter(
        is_staff=False
    ).count()

    transactions_count = Transaction.objects.count()

    budgets_count = Budget.objects.count()

    categories_count = Category.objects.filter(
        is_active=True
    ).count()

    return render(
        request,
        'tracker/admin_dashboard.html',
        {
            'users_count': users_count,
            'transactions_count': transactions_count,
            'budgets_count': budgets_count,
            'categories_count': categories_count,
        }
    )

def user_logout(request):
    logout(request)
    return redirect('login')


# =========================================================
# RECURRING TRANSACTIONS
# =========================================================

def calculate_next_recurring_date(
    start_date,
    frequency,
    day_of_month
):
    """
    Calculate the first scheduled occurrence for a recurring
    transaction.
    """

    if frequency == 'yearly':
        target_year = start_date.year

        if day_of_month <= start_date.day:
            target_year += 1

        last_day = calendar.monthrange(
            target_year,
            start_date.month
        )[1]

        return date(
            target_year,
            start_date.month,
            min(day_of_month, last_day)
        )

    # Monthly frequency.
    target_year = start_date.year
    target_month = start_date.month

    if day_of_month <= start_date.day:
        if target_month == 12:
            target_year += 1
            target_month = 1
        else:
            target_month += 1

    last_day = calendar.monthrange(
        target_year,
        target_month
    )[1]

    return date(
        target_year,
        target_month,
        min(day_of_month, last_day)
    )


@login_required
def recurring_transaction_list(request):
    process_due_recurring_transactions(
        request.user
    )

    recurring_transactions = (
        request.user.recurring_transactions
        .select_related('category')
        .order_by(
            '-is_active',
            'next_date',
            '-created_at'
        )
    )

    active_count = recurring_transactions.filter(
        is_active=True
    ).count()

    inactive_count = recurring_transactions.filter(
        is_active=False
    ).count()

    return render(
        request,
        'tracker/recurring_transactions.html',
        {
            'recurring_transactions': recurring_transactions,
            'active_count': active_count,
            'inactive_count': inactive_count,
        }
    )


@login_required
def add_recurring_transaction(request):
    if request.method == 'POST':
        form = RecurringTransactionForm(
            request.POST
        )

        if form.is_valid():
            recurring_transaction = form.save(
                commit=False
            )

            recurring_transaction.user = request.user

            recurring_transaction.next_date = (
                calculate_next_recurring_date(
                    recurring_transaction.start_date,
                    recurring_transaction.frequency,
                    recurring_transaction.day_of_month
                )
            )

            recurring_transaction.save()

            messages.success(
                request,
                'Recurring transaction created successfully.'
            )

            return redirect(
                'recurring_transaction_list'
            )

    else:
        form = RecurringTransactionForm()

    return render(
        request,
        'tracker/add_recurring_transaction.html',
        {
            'form': form,
        }
    )




@login_required
def edit_recurring_transaction(
    request,
    recurring_id
):
    recurring_transaction = get_object_or_404(
        request.user.recurring_transactions,
        id=recurring_id
    )

    if request.method == 'POST':
        form = RecurringTransactionForm(
            request.POST,
            instance=recurring_transaction
        )

        if form.is_valid():
            recurring_transaction = form.save(
                commit=False
            )

            # Recalculate the next scheduled occurrence
            # whenever the schedule is edited.
            recurring_transaction.next_date = (
                calculate_next_recurring_date(
                    recurring_transaction.start_date,
                    recurring_transaction.frequency,
                    recurring_transaction.day_of_month
                )
            )

            recurring_transaction.save()

            messages.success(
                request,
                'Recurring transaction updated successfully.'
            )

            return redirect(
                'recurring_transaction_list'
            )

    else:
        form = RecurringTransactionForm(
            instance=recurring_transaction
        )

    return render(
        request,
        'tracker/edit_recurring_transaction.html',
        {
            'form': form,
            'recurring_transaction': recurring_transaction,
        }
    )


@login_required
def toggle_recurring_transaction(
    request,
    recurring_id
):
    recurring_transaction = get_object_or_404(
        request.user.recurring_transactions,
        id=recurring_id
    )

    if request.method != 'POST':
        return redirect(
            'recurring_transaction_list'
        )

    recurring_transaction.is_active = (
        not recurring_transaction.is_active
    )

    recurring_transaction.save(
        update_fields=[
            'is_active',
            'updated_at',
        ]
    )

    if recurring_transaction.is_active:
        messages.success(
            request,
            'Recurring transaction activated.'
        )
    else:
        messages.success(
            request,
            'Recurring transaction paused.'
        )

    return redirect(
        'recurring_transaction_list'
    )


@login_required
def delete_recurring_transaction(
    request,
    recurring_id
):
    recurring_transaction = get_object_or_404(
        request.user.recurring_transactions,
        id=recurring_id
    )

    if request.method == 'POST':
        recurring_transaction.delete()

        messages.success(
            request,
            'Recurring transaction deleted successfully.'
        )

    return redirect(
        'recurring_transaction_list'
    )



@login_required
def add_transaction(request):
    if request.method == 'POST':
        form = TransactionForm(request.POST)

        if form.is_valid():
            transaction = form.save(commit=False)
            transaction.user = request.user
            transaction.save()

            messages.success(
                request,
                'Transaction added successfully.'
            )

            return redirect('dashboard')
    else:
        form = TransactionForm()

    return render(
        request,
        'tracker/add_transaction.html',
        {'form': form}
    )

@login_required
def transaction_list(request):
    transactions = request.user.transactions.select_related(
        'category'
    ).order_by('-date', '-created_at')

    search_query = request.GET.get('search', '').strip()
    category_filter = request.GET.get('category', '').strip()
    type_filter = request.GET.get('type', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()

    if search_query:
        transactions = transactions.filter(
            models.Q(description__icontains=search_query)
            | models.Q(category__name__icontains=search_query)
            | models.Q(payment_method__icontains=search_query)
        )
    if category_filter:
        transactions = transactions.filter(
            category_id=category_filter
        )
    if type_filter:
        transactions = transactions.filter(
            type=type_filter
        )
    if date_from:
        transactions = transactions.filter(
            date__gte=date_from
        )

    if date_to:
        transactions = transactions.filter(
            date__lte=date_to
        )
    

    return render(
        request,
        'tracker/transaction_list.html',
        {
            'transactions': transactions,
            'search_query': search_query,
            'categories': Category.objects.filter(is_active=True),
            'category_filter': category_filter,
            'type_filter': type_filter,
            'date_from': date_from,
            'date_to': date_to,
        }
    )

@login_required
def edit_transaction(request, transaction_id):
    transaction = get_object_or_404(
        request.user.transactions,
        id=transaction_id
    )

    if request.method == 'POST':
        form = TransactionForm(
            request.POST,
            instance=transaction
        )

        if form.is_valid():
            form.save()

            messages.success(
                request,
                'Transaction updated successfully.'
            )

            return redirect('transaction_list')
    else:
        form = TransactionForm(instance=transaction)

    return render(
        request,
        'tracker/edit_transaction.html',
        {
            'form': form,
            'transaction': transaction,
        }
    )

@login_required
def delete_transaction(request, transaction_id):
    transaction = get_object_or_404(
        request.user.transactions,
        id=transaction_id
    )

    if request.method == 'POST':
        transaction.delete()

        messages.success(
            request,
            'Transaction deleted successfully.'
        )

        return redirect('transaction_list')

    return render(
        request,
        'tracker/delete_transaction.html',
        {
            'transaction': transaction,
        }
    )

@login_required
def add_budget(request):
    if request.method == 'POST':
        form = BudgetForm(request.POST)

        if form.is_valid():
            budget = form.save(commit=False)
            budget.user = request.user
            budget.save()

            messages.success(
                request,
                'Budget created successfully.'
            )

            return redirect('dashboard')
    else:
        form = BudgetForm()

    return render(
        request,
        'tracker/add_budget.html',
        {'form': form}
    )

@login_required
def budget_list(request):
    budgets = request.user.budgets.select_related(
        'category'
    ).order_by('-start_date')

    for budget in budgets:
        spent = Transaction.objects.filter(
            user=request.user,
            category=budget.category,
            type='expense',
            date__gte=budget.start_date,
            date__lte=budget.end_date,
        ).aggregate(
            total=Sum('amount')
        )['total'] or 0

        budget.spent = spent
        budget.remaining = budget.limit - spent
        budget.is_exceeded = spent >= budget.limit

    return render(
        request,
        'tracker/budget_list.html',
        {
            'budgets': budgets,
        }
    )

@login_required
def edit_budget(request, budget_id):
    budget = get_object_or_404(
        request.user.budgets,
        id=budget_id
    )

    if request.method == 'POST':
        form = BudgetForm(
            request.POST,
            instance=budget
        )

        if form.is_valid():
            form.save()

            messages.success(
                request,
                'Budget updated successfully.'
            )

            return redirect('budget_list')
    else:
        form = BudgetForm(instance=budget)

    return render(
        request,
        'tracker/edit_budget.html',
        {
            'form': form,
            'budget': budget,
        }
    )

@login_required
def report_view(request):
    """
    Build the interactive Reports & Analytics page.

    NOTE:
    This view powers only report.html.
    export_report_pdf() and export_report_excel() remain untouched.
    """
    from calendar import monthrange
    from datetime import timedelta
    from decimal import Decimal

    report_type = request.GET.get(
        'report_type',
        'monthly'
    )

    if report_type not in ('monthly', 'yearly'):
        report_type = 'monthly'

    today = timezone.localdate()

    # =========================================================
    # SELECTED PERIOD
    # =========================================================

    if report_type == 'yearly':
        start_date = today.replace(
            month=1,
            day=1
        )

        end_date = today.replace(
            month=12,
            day=31
        )

        previous_start_date = start_date.replace(
            year=start_date.year - 1
        )

        previous_end_date = end_date.replace(
            year=end_date.year - 1
        )

    else:
        start_date = today.replace(
            day=1
        )

        if today.month == 12:
            next_month = today.replace(
                year=today.year + 1,
                month=1,
                day=1
            )
        else:
            next_month = today.replace(
                month=today.month + 1,
                day=1
            )

        end_date = next_month - timedelta(
            days=1
        )

        if start_date.month == 1:
            previous_start_date = start_date.replace(
                year=start_date.year - 1,
                month=12
            )
        else:
            previous_start_date = start_date.replace(
                month=start_date.month - 1
            )

        previous_end_date = start_date - timedelta(
            days=1
        )

    # =========================================================
    # CURRENT + PREVIOUS TRANSACTIONS
    # =========================================================

    transactions = request.user.transactions.filter(
        date__gte=start_date,
        date__lte=end_date
    ).select_related(
        'category'
    )

    previous_transactions = request.user.transactions.filter(
        date__gte=previous_start_date,
        date__lte=previous_end_date
    )

    # =========================================================
    # CORE FINANCIAL TOTALS
    # =========================================================

    total_income = transactions.filter(
        type='income'
    ).aggregate(
        total=Sum('amount')
    )['total'] or Decimal('0')

    total_expenses = transactions.filter(
        type='expense'
    ).aggregate(
        total=Sum('amount')
    )['total'] or Decimal('0')

    balance = total_income - total_expenses

    transaction_count = transactions.count()

    income_transaction_count = transactions.filter(
        type='income'
    ).count()

    expense_transaction_count = transactions.filter(
        type='expense'
    ).count()

    # =========================================================
    # KEY FINANCIAL METRICS
    # =========================================================

    if total_income > 0:
        savings_rate = (
            balance / total_income
        ) * Decimal('100')

        expense_ratio = (
            total_expenses / total_income
        ) * Decimal('100')
    else:
        savings_rate = Decimal('0')
        expense_ratio = Decimal('0')

    days_in_period = (
        end_date - start_date
    ).days + 1

    if days_in_period > 0:
        average_daily_expense = (
            total_expenses / Decimal(
                str(days_in_period)
            )
        )
    else:
        average_daily_expense = Decimal('0')

    if expense_transaction_count > 0:
        average_expense_transaction = (
            total_expenses / Decimal(
                str(expense_transaction_count)
            )
        )
    else:
        average_expense_transaction = Decimal('0')

    savings_rate = round(
        savings_rate,
        1
    )

    expense_ratio = round(
        expense_ratio,
        1
    )

    average_daily_expense = round(
        average_daily_expense,
        2
    )

    average_expense_transaction = round(
        average_expense_transaction,
        2
    )

    # =========================================================
    # PREVIOUS-PERIOD COMPARISON
    # =========================================================

    previous_income = previous_transactions.filter(
        type='income'
    ).aggregate(
        total=Sum('amount')
    )['total'] or Decimal('0')

    previous_expenses = previous_transactions.filter(
        type='expense'
    ).aggregate(
        total=Sum('amount')
    )['total'] or Decimal('0')

    previous_balance = (
        previous_income - previous_expenses
    )

    def percentage_change(current, previous):
        if previous == 0:
            if current == 0:
                return Decimal('0')
            return Decimal('100')

        return round(
            (
                (current - previous)
                / previous
            ) * Decimal('100'),
            1
        )

    income_change = percentage_change(
        total_income,
        previous_income
    )

    expense_change = percentage_change(
        total_expenses,
        previous_expenses
    )

    balance_change = percentage_change(
        balance,
        previous_balance
    )

    # =========================================================
    # CATEGORY ANALYSIS
    # =========================================================

    expense_categories = list(
        transactions.filter(
            type='expense'
        ).values(
            'category__name'
        ).annotate(
            total=Sum('amount')
        ).order_by(
            '-total'
        )
    )

    income_categories = list(
        transactions.filter(
            type='income'
        ).values(
            'category__name'
        ).annotate(
            total=Sum('amount')
        ).order_by(
            '-total'
        )
    )

    expense_category_summary = []

    for item in expense_categories:
        total = item['total'] or Decimal('0')

        if total_expenses > 0:
            percentage = round(
                (
                    total / total_expenses
                ) * Decimal('100'),
                1
            )
        else:
            percentage = Decimal('0')

        count = transactions.filter(
            type='expense',
            category__name=item['category__name']
        ).count()

        expense_category_summary.append({
            'name': item['category__name'],
            'total': total,
            'percentage': percentage,
            'count': count,
        })

    income_source_summary = []

    for item in income_categories:
        total = item['total'] or Decimal('0')

        if total_income > 0:
            percentage = round(
                (
                    total / total_income
                ) * Decimal('100'),
                1
            )
        else:
            percentage = Decimal('0')

        count = transactions.filter(
            type='income',
            category__name=item['category__name']
        ).count()

        income_source_summary.append({
            'name': item['category__name'],
            'total': total,
            'percentage': percentage,
            'count': count,
        })

    # =========================================================
    # HIGHLIGHTS
    # =========================================================

    highest_expense = transactions.filter(
        type='expense'
    ).order_by(
        '-amount'
    ).first()

    highest_income = transactions.filter(
        type='income'
    ).order_by(
        '-amount'
    ).first()

    top_expense_category = (
        expense_category_summary[0]
        if expense_category_summary
        else None
    )

    top_income_source = (
        income_source_summary[0]
        if income_source_summary
        else None
    )

    # =========================================================
    # PAYMENT-METHOD ANALYSIS
    # =========================================================

    payment_method_totals = list(
        transactions.values(
            'payment_method'
        ).annotate(
            total=Sum('amount')
        ).order_by(
            '-total'
        )
    )

    payment_method_labels = {
        'cash': 'Cash',
        'upi': 'UPI / Wallet',
        'card': 'Card',
    }

    payment_summary = []

    for item in payment_method_totals:
        amount = item['total'] or Decimal('0')
        method = item['payment_method']

        count = transactions.filter(
            payment_method=method
        ).count()

        if total_income + total_expenses > 0:
            percentage = round(
                (
                    amount
                    / (
                        total_income
                        + total_expenses
                    )
                ) * Decimal('100'),
                1
            )
        else:
            percentage = Decimal('0')

        payment_summary.append({
            'method': payment_method_labels.get(
                method,
                method.title()
            ),
            'code': method,
            'amount': amount,
            'count': count,
            'percentage': percentage,
        })

    if payment_summary:
        most_used_payment_method = (
            payment_summary[0]['method']
        )
    else:
        most_used_payment_method = '—'

    # =========================================================
    # MONTHLY / DAILY TREND
    # =========================================================

    trend_labels = []
    trend_income = []
    trend_expenses = []
    trend_savings = []

    if report_type == 'yearly':
        current_month = start_date

        while current_month <= end_date:
            if current_month.month == 12:
                next_period = current_month.replace(
                    year=current_month.year + 1,
                    month=1,
                    day=1
                )
            else:
                next_period = current_month.replace(
                    month=current_month.month + 1,
                    day=1
                )

            period_end = (
                next_period - timedelta(days=1)
            )

            month_transactions = transactions.filter(
                date__gte=current_month,
                date__lte=period_end
            )

            month_income = (
                month_transactions.filter(
                    type='income'
                ).aggregate(
                    total=Sum('amount')
                )['total'] or Decimal('0')
            )

            month_expenses = (
                month_transactions.filter(
                    type='expense'
                ).aggregate(
                    total=Sum('amount')
                )['total'] or Decimal('0')
            )

            month_savings = (
                month_income - month_expenses
            )

            trend_labels.append(
                current_month.strftime('%b')
            )

            trend_income.append(
                float(month_income)
            )

            trend_expenses.append(
                float(month_expenses)
            )

            trend_savings.append(
                float(month_savings)
            )

            current_month = next_period

    else:
        current_day = start_date

        while current_day <= end_date:
            day_transactions = transactions.filter(
                date=current_day
            )

            day_income = (
                day_transactions.filter(
                    type='income'
                ).aggregate(
                    total=Sum('amount')
                )['total'] or Decimal('0')
            )

            day_expenses = (
                day_transactions.filter(
                    type='expense'
                ).aggregate(
                    total=Sum('amount')
                )['total'] or Decimal('0')
            )

            day_savings = (
                day_income - day_expenses
            )

            trend_labels.append(
                current_day.strftime('%d')
            )

            trend_income.append(
                float(day_income)
            )

            trend_expenses.append(
                float(day_expenses)
            )

            trend_savings.append(
                float(day_savings)
            )

            current_day += timedelta(days=1)

    # =========================================================
    # BUDGET ANALYSIS
    # =========================================================

    budgets = request.user.budgets.filter(
        start_date__lte=end_date,
        end_date__gte=start_date
    ).select_related(
        'category'
    ).order_by(
        'category__name'
    )

    budget_analysis = []

    for budget in budgets:
        spent = transactions.filter(
            category=budget.category,
            type='expense',
            date__gte=budget.start_date,
            date__lte=budget.end_date
        ).aggregate(
            total=Sum('amount')
        )['total'] or Decimal('0')

        limit = budget.limit or Decimal('0')

        if limit > 0:
            utilization = round(
                (
                    spent / limit
                ) * Decimal('100'),
                1
            )
        else:
            utilization = Decimal('0')

        if utilization >= 100:
            status = 'Exceeded'
        elif utilization >= 80:
            status = 'Near Limit'
        else:
            status = 'On Track'

        budget_analysis.append({
            'category': budget.category.name,
            'limit': limit,
            'spent': spent,
            'remaining': limit - spent,
            'utilization': utilization,
            'status': status,
            'period': budget.period.title(),
        })

    total_budget_limit = sum(
        item['limit']
        for item in budget_analysis
    )

    total_budget_spent = sum(
        item['spent']
        for item in budget_analysis
    )

    if total_budget_limit > 0:
        overall_budget_utilization = round(
            (
                total_budget_spent
                / total_budget_limit
            ) * Decimal('100'),
            1
        )
    else:
        overall_budget_utilization = Decimal('0')

    exceeded_budget_count = sum(
        1
        for item in budget_analysis
        if item['status'] == 'Exceeded'
    )

    near_limit_budget_count = sum(
        1
        for item in budget_analysis
        if item['status'] == 'Near Limit'
    )

    # =========================================================
    # REPORT CONTEXT
    # =========================================================

    return render(
        request,
        'tracker/report.html',
        {
            # Existing fields — preserved.
            'report_type': report_type,
            'start_date': start_date,
            'end_date': end_date,
            'total_income': total_income,
            'total_expenses': total_expenses,
            'balance': balance,
            'category_summary': (
                transactions.values(
                    'category__name',
                    'category__type'
                ).annotate(
                    total=Sum('amount')
                ).order_by(
                    '-total'
                )
            ),

            # Overview.
            'transaction_count': transaction_count,
            'income_transaction_count': income_transaction_count,
            'expense_transaction_count': expense_transaction_count,
            'savings_rate': savings_rate,
            'expense_ratio': expense_ratio,
            'average_daily_expense': average_daily_expense,
            'average_expense_transaction': average_expense_transaction,

            # Comparison.
            'previous_start_date': previous_start_date,
            'previous_end_date': previous_end_date,
            'previous_income': previous_income,
            'previous_expenses': previous_expenses,
            'previous_balance': previous_balance,
            'income_change': income_change,
            'expense_change': expense_change,
            'balance_change': balance_change,

            # Highlights.
            'highest_expense': highest_expense,
            'highest_income': highest_income,
            'top_expense_category': top_expense_category,
            'top_income_source': top_income_source,
            'most_used_payment_method': most_used_payment_method,

            # Category + payment breakdowns.
            'expense_category_summary': expense_category_summary,
            'income_source_summary': income_source_summary,
            'payment_summary': payment_summary,

            # Trends for Chart.js.
            'trend_labels': trend_labels,
            'trend_income': trend_income,
            'trend_expenses': trend_expenses,
            'trend_savings': trend_savings,

            # Budget analysis.
            'budget_analysis': budget_analysis,
            'total_budget_limit': total_budget_limit,
            'total_budget_spent': total_budget_spent,
            'overall_budget_utilization': overall_budget_utilization,
            'exceeded_budget_count': exceeded_budget_count,
            'near_limit_budget_count': near_limit_budget_count,
        }
    )

@login_required
def export_report_pdf(request):
    report_type = request.GET.get('report_type', 'monthly')

    today = timezone.localdate()

    if report_type == 'yearly':
        start_date = today.replace(
            month=1,
            day=1
        )

        end_date = today.replace(
            month=12,
            day=31
        )

    else:
        start_date = today.replace(
            day=1
        )

        if today.month == 12:
            next_month = today.replace(
                year=today.year + 1,
                month=1,
                day=1
            )

        else:
            next_month = today.replace(
                month=today.month + 1,
                day=1
            )

        end_date = next_month - timedelta(days=1)


    # =========================================================
    # FETCH REPORT DATA
    # =========================================================

    transactions = request.user.transactions.filter(
        date__gte=start_date,
        date__lte=end_date
    )


    total_income = transactions.filter(
        type='income'
    ).aggregate(
        total=Sum('amount')
    )['total'] or 0


    total_expenses = transactions.filter(
        type='expense'
    ).aggregate(
        total=Sum('amount')
    )['total'] or 0


    balance = total_income - total_expenses


    transaction_count = transactions.count()


    category_summary = transactions.values(
        'category__name',
        'category__type'
    ).annotate(
        total=Sum('amount')
    ).order_by('-total')


    # =========================================================
    # PDF RESPONSE
    # =========================================================

    response = HttpResponse(
        content_type='application/pdf'
    )

    filename = (
        f'Expense_Tracker_'
        f'{report_type.title()}_Report.pdf'
    )

    response['Content-Disposition'] = (
        f'attachment; filename="{filename}"'
    )


    # =========================================================
    # DOCUMENT
    # =========================================================

    document = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=42,
        bottomMargin=42,
        title=(
            f'Expense Tracker - '
            f'{report_type.title()} Report'
        ),
        author='Expense Tracker',
    )


    # =========================================================
    # COLORS
    # =========================================================

    INDIGO = colors.HexColor('#4F46E5')
    INDIGO_DARK = colors.HexColor('#3730A3')
    INDIGO_LIGHT = colors.HexColor('#EEF2FF')

    TEXT_DARK = colors.HexColor('#111827')
    TEXT_MUTED = colors.HexColor('#6B7280')

    BORDER = colors.HexColor('#E5E7EB')
    ROW_LIGHT = colors.HexColor('#F9FAFB')

    GREEN = colors.HexColor('#059669')
    GREEN_LIGHT = colors.HexColor('#ECFDF5')

    RED = colors.HexColor('#DC2626')
    RED_LIGHT = colors.HexColor('#FEF2F2')

    BLUE_LIGHT = colors.HexColor('#EFF6FF')


    # =========================================================
    # STYLES
    # =========================================================

    styles = getSampleStyleSheet()


    report_title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.white,
        alignment=TA_LEFT,
        spaceAfter=0,
    )


    report_type_style = ParagraphStyle(
        'ReportType',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.white,
        alignment=TA_RIGHT,
    )


    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=TEXT_MUTED,
    )


    section_style = ParagraphStyle(
        'Section',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=TEXT_DARK,
        spaceAfter=4,
    )


    small_style = ParagraphStyle(
        'Small',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=TEXT_MUTED,
    )


    card_label_style = ParagraphStyle(
        'CardLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=TEXT_MUTED,
        alignment=TA_CENTER,
    )


    income_value_style = ParagraphStyle(
        'IncomeValue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=GREEN,
        alignment=TA_CENTER,
    )


    expense_value_style = ParagraphStyle(
        'ExpenseValue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=RED,
        alignment=TA_CENTER,
    )


    balance_value_style = ParagraphStyle(
        'BalanceValue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=INDIGO,
        alignment=TA_CENTER,
    )


    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=TA_LEFT,
    )


    table_text_style = ParagraphStyle(
        'TableText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=TEXT_DARK,
    )


    table_amount_style = ParagraphStyle(
        'TableAmount',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        alignment=TA_RIGHT,
    )


    # =========================================================
    # PAGE DECORATION
    # =========================================================

    def draw_page_decor(canvas, doc):
        canvas.saveState()

        page_width, page_height = A4

        # Top accent
        canvas.setFillColor(INDIGO)
        canvas.rect(
            0,
            page_height - 7,
            page_width,
            7,
            stroke=0,
            fill=1
        )

        # Footer line
        canvas.setStrokeColor(BORDER)
        canvas.setLineWidth(0.6)
        canvas.line(
            36,
            30,
            page_width - 36,
            30
        )

        # Footer text
        canvas.setFillColor(TEXT_MUTED)
        canvas.setFont(
            'Helvetica',
            7.5
        )

        canvas.drawString(
            36,
            18,
            'Expense Tracker • Financial Report'
        )

        canvas.drawRightString(
            page_width - 36,
            18,
            f'Page {doc.page}'
        )

        canvas.restoreState()


    # =========================================================
    # CONTENT
    # =========================================================

    elements = []


    # =========================================================
    # HEADER
    # =========================================================

    user_name = (
        request.user.get_full_name().strip()
        or request.user.username
    )


    header_table = Table(
        [
            [
                Paragraph(
                    'EXPENSE TRACKER',
                    report_title_style
                ),

                Paragraph(
                    f'{report_type.title()} Report',
                    report_type_style
                )
            ]
        ],
        colWidths=[
            350,
            173
        ],
        rowHeights=[70],
    )


    header_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (-1, -1),
                INDIGO
            ),

            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),

            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                18
            ),

            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                18
            ),

            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                10
            ),

            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                10
            ),
        ])
    )


    elements.append(header_table)

    elements.append(
        Spacer(1, 14)
    )


    # =========================================================
    # REPORT INFORMATION
    # =========================================================

    info_table = Table(
        [
            [
                Paragraph(
                    f'<b>Prepared for:</b> '
                    f'{escape(user_name)}',
                    small_style
                ),

                Paragraph(
                    f'<b>Period:</b> '
                    f'{start_date.strftime("%d %b %Y")} '
                    f'– '
                    f'{end_date.strftime("%d %b %Y")}',
                    small_style
                )
            ],

            [
                Paragraph(
                    f'<b>Generated:</b> '
                    f'{today.strftime("%d %b %Y")}',
                    small_style
                ),

                Paragraph(
                    f'<b>Transactions:</b> '
                    f'{transaction_count}',
                    small_style
                )
            ]
        ],
        colWidths=[
            261.5,
            261.5
        ],
    )


    info_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (-1, -1),
                colors.white
            ),

            (
                'BOX',
                (0, 0),
                (-1, -1),
                0.7,
                BORDER
            ),

            (
                'INNERGRID',
                (0, 0),
                (-1, -1),
                0.4,
                BORDER
            ),

            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                12
            ),

            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                12
            ),

            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                8
            ),

            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )


    elements.append(info_table)

    elements.append(
        Spacer(1, 18)
    )


    # =========================================================
    # FINANCIAL SUMMARY
    # =========================================================

    elements.append(
        Paragraph(
            'Financial Overview',
            section_style
        )
    )

    elements.append(
        Paragraph(
            'A summary of your financial activity for the selected period.',
            subtitle_style
        )
    )

    elements.append(
        Spacer(1, 9)
    )


    income_text = (
        f'Rs. {total_income:,.2f}'
    )

    expense_text = (
        f'Rs. {total_expenses:,.2f}'
    )

    balance_text = (
        f'Rs. {balance:,.2f}'
    )


    summary_table = Table(
        [
            [
                [
                    Paragraph(
                        'TOTAL INCOME',
                        card_label_style
                    ),

                    Spacer(1, 6),

                    Paragraph(
                        income_text,
                        income_value_style
                    )
                ],

                [
                    Paragraph(
                        'TOTAL EXPENSES',
                        card_label_style
                    ),

                    Spacer(1, 6),

                    Paragraph(
                        expense_text,
                        expense_value_style
                    )
                ],

                [
                    Paragraph(
                        'CURRENT BALANCE',
                        card_label_style
                    ),

                    Spacer(1, 6),

                    Paragraph(
                        balance_text,
                        balance_value_style
                    )
                ]
            ]
        ],
        colWidths=[
            174,
            174,
            174
        ],
    )


    summary_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (0, 0),
                GREEN_LIGHT
            ),

            (
                'BACKGROUND',
                (1, 0),
                (1, 0),
                RED_LIGHT
            ),

            (
                'BACKGROUND',
                (2, 0),
                (2, 0),
                INDIGO_LIGHT
            ),

            (
                'BOX',
                (0, 0),
                (-1, -1),
                0.7,
                BORDER
            ),

            (
                'INNERGRID',
                (0, 0),
                (-1, -1),
                0.7,
                BORDER
            ),

            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),

            (
                'ALIGN',
                (0, 0),
                (-1, -1),
                'CENTER'
            ),

            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                14
            ),

            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                14
            ),

            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                8
            ),

            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )


    elements.append(summary_table)

    elements.append(
        Spacer(1, 22)
    )


    # =========================================================
    # CATEGORY SUMMARY
    # =========================================================

    elements.append(
        Paragraph(
            'Category-wise Summary',
            section_style
        )
    )

    elements.append(
        Paragraph(
            'Breakdown of income and expenses by category.',
            subtitle_style
        )
    )

    elements.append(
        Spacer(1, 9)
    )


    table_data = [
        [
            Paragraph(
                'CATEGORY',
                table_header_style
            ),

            Paragraph(
                'TYPE',
                table_header_style
            ),

            Paragraph(
                'TOTAL',
                table_header_style
            )
        ]
    ]


    for item in category_summary:

        category_name = escape(
            str(item['category__name'])
        )

        category_type = item[
            'category__type'
        ].title()

        total_value = item['total']


        if item['category__type'] == 'income':

            type_color = '#059669'

        else:

            type_color = '#DC2626'


        table_data.append(
            [
                Paragraph(
                    category_name,
                    table_text_style
                ),

                Paragraph(
                    (
                        f'<font color="{type_color}">'
                        f'<b>{category_type}</b>'
                        f'</font>'
                    ),
                    table_text_style
                ),

                Paragraph(
                    f'Rs. {total_value:,.2f}',
                    table_amount_style
                )
            ]
        )


    if len(table_data) > 1:

        category_table = Table(
            table_data,
            colWidths=[
                230,
                105,
                188
            ],
            repeatRows=1,
        )


        category_table_style = [
            (
                'BACKGROUND',
                (0, 0),
                (-1, 0),
                INDIGO
            ),

            (
                'TEXTCOLOR',
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                'BOX',
                (0, 0),
                (-1, -1),
                0.7,
                BORDER
            ),

            (
                'INNERGRID',
                (0, 0),
                (-1, -1),
                0.4,
                BORDER
            ),

            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),

            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                10
            ),

            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                10
            ),

            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                8
            ),

            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                8
            ),
        ]


        # Alternating row backgrounds
        for row_index in range(
            1,
            len(table_data)
        ):

            if row_index % 2 == 0:

                category_table_style.append(
                    (
                        'BACKGROUND',
                        (0, row_index),
                        (-1, row_index),
                        ROW_LIGHT
                    )
                )


        category_table.setStyle(
            TableStyle(category_table_style)
        )


        elements.append(category_table)


    else:

        empty_table = Table(
            [
                [
                    Paragraph(
                        'No transactions found for this period.',
                        small_style
                    )
                ]
            ],
            colWidths=[523],
        )


        empty_table.setStyle(
            TableStyle([
                (
                    'BACKGROUND',
                    (0, 0),
                    (-1, -1),
                    BLUE_LIGHT
                ),

                (
                    'BOX',
                    (0, 0),
                    (-1, -1),
                    0.7,
                    BORDER
                ),

                (
                    'LEFTPADDING',
                    (0, 0),
                    (-1, -1),
                    12
                ),

                (
                    'RIGHTPADDING',
                    (0, 0),
                    (-1, -1),
                    12
                ),

                (
                    'TOPPADDING',
                    (0, 0),
                    (-1, -1),
                    14
                ),

                (
                    'BOTTOMPADDING',
                    (0, 0),
                    (-1, -1),
                    14
                ),
            ])
        )


        elements.append(empty_table)


    elements.append(
        Spacer(1, 20)
    )


    # =========================================================
    # REPORT NOTE
    # =========================================================

    note_table = Table(
        [
            [
                Paragraph(
                    '<b>Expense Tracker</b><br/>'
                    'This report is generated from the transactions '
                    'recorded in your account for the selected period.',
                    small_style
                )
            ]
        ],
        colWidths=[523],
    )


    note_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (-1, -1),
                INDIGO_LIGHT
            ),

            (
                'BOX',
                (0, 0),
                (-1, -1),
                0.7,
                colors.HexColor('#C7D2FE')
            ),

            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                12
            ),

            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                12
            ),

            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                10
            ),

            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                10
            ),
        ])
    )


    elements.append(note_table)


    # =========================================================
    # BUILD PDF
    # =========================================================

    document.build(
        elements,
        onFirstPage=draw_page_decor,
        onLaterPages=draw_page_decor
    )


    # =========================================================
    # SAVE REPORT HISTORY
    # =========================================================

    Report.objects.create(
        user=request.user,
        report_type=report_type,
        export_format='pdf',
        start_date=start_date,
        end_date=end_date
    )


    return response

@login_required
def export_report_excel(request):
    report_type = request.GET.get('report_type', 'monthly')

    today = timezone.localdate()

    if report_type == 'yearly':
        start_date = today.replace(
            month=1,
            day=1
        )

        end_date = today.replace(
            month=12,
            day=31
        )

    else:
        start_date = today.replace(
            day=1
        )

        if today.month == 12:
            next_month = today.replace(
                year=today.year + 1,
                month=1,
                day=1
            )

        else:
            next_month = today.replace(
                month=today.month + 1,
                day=1
            )

        end_date = next_month - timedelta(days=1)


    # =========================================================
    # FETCH REPORT DATA
    # =========================================================

    transactions = request.user.transactions.filter(
        date__gte=start_date,
        date__lte=end_date
    )


    total_income = transactions.filter(
        type='income'
    ).aggregate(
        total=Sum('amount')
    )['total'] or 0


    total_expenses = transactions.filter(
        type='expense'
    ).aggregate(
        total=Sum('amount')
    )['total'] or 0


    balance = total_income - total_expenses


    transaction_count = transactions.count()


    category_summary = transactions.values(
        'category__name',
        'category__type'
    ).annotate(
        total=Sum('amount')
    ).order_by('-total')


    # =========================================================
    # WORKBOOK
    # =========================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = 'Financial Report'

    worksheet.sheet_view.showGridLines = False


    # =========================================================
    # COLORS
    # =========================================================

    INDIGO = '4F46E5'
    INDIGO_DARK = '3730A3'
    INDIGO_LIGHT = 'EEF2FF'

    GREEN = '059669'
    GREEN_LIGHT = 'ECFDF5'

    RED = 'DC2626'
    RED_LIGHT = 'FEF2F2'

    TEXT_DARK = '111827'
    TEXT_MUTED = '6B7280'

    BORDER_COLOR = 'E5E7EB'
    ROW_LIGHT = 'F9FAFB'

    WHITE = 'FFFFFF'


    # =========================================================
    # REUSABLE STYLES
    # =========================================================

    thin_side = Side(
        style='thin',
        color=BORDER_COLOR
    )


    medium_side = Side(
        style='medium',
        color=INDIGO
    )


    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )


    # =========================================================
    # PAGE SETTINGS
    # =========================================================

    worksheet.freeze_panes = 'A10'

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.page_setup.fitToWidth = 1

    worksheet.page_setup.fitToHeight = 0

    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5


    # =========================================================
    # HEADER
    # =========================================================

    worksheet.merge_cells('A1:C2')

    title_cell = worksheet['A1']

    title_cell.value = (
        f'EXPENSE TRACKER\n'
        f'{report_type.title()} Financial Report'
    )

    title_cell.font = Font(
        name='Calibri',
        size=18,
        bold=True,
        color=WHITE
    )

    title_cell.fill = PatternFill(
        'solid',
        fgColor=INDIGO
    )

    title_cell.alignment = Alignment(
        horizontal='left',
        vertical='center',
        wrap_text=True
    )


    for row in worksheet['A1:C2']:

        for cell in row:

            cell.fill = PatternFill(
                'solid',
                fgColor=INDIGO
            )


    worksheet.row_dimensions[1].height = 28
    worksheet.row_dimensions[2].height = 28


    # =========================================================
    # REPORT INFORMATION
    # =========================================================

    worksheet['A4'] = 'Report Information'

    worksheet['A4'].font = Font(
        bold=True,
        size=13,
        color=INDIGO_DARK
    )


    worksheet['A5'] = 'Prepared For'

    worksheet['B5'] = (
        request.user.get_full_name().strip()
        or request.user.username
    )


    worksheet['A6'] = 'Period'

    worksheet['B6'] = (
        f'{start_date.strftime("%d %b %Y")} '
        f'to '
        f'{end_date.strftime("%d %b %Y")}'
    )


    worksheet['A7'] = 'Generated'

    worksheet['B7'] = today.strftime(
        '%d %b %Y'
    )


    worksheet['A8'] = 'Transactions'

    worksheet['B8'] = transaction_count


    for row in range(5, 9):

        worksheet[f'A{row}'].font = Font(
            bold=True,
            color=TEXT_DARK
        )

        worksheet[f'B{row}'].font = Font(
            color=TEXT_MUTED
        )

        worksheet[f'A{row}'].fill = PatternFill(
            'solid',
            fgColor=INDIGO_LIGHT
        )

        worksheet[f'B{row}'].fill = PatternFill(
            'solid',
            fgColor=WHITE
        )

        worksheet[f'A{row}'].border = thin_border
        worksheet[f'B{row}'].border = thin_border


    # =========================================================
    # FINANCIAL SUMMARY
    # =========================================================

    worksheet['A10'] = 'Financial Overview'

    worksheet['A10'].font = Font(
        bold=True,
        size=13,
        color=INDIGO_DARK
    )


    summary_headers = [
        'Total Income',
        'Total Expenses',
        'Current Balance'
    ]


    summary_values = [
        float(total_income),
        float(total_expenses),
        float(balance)
    ]


    for column, (header, value) in enumerate(
        zip(summary_headers, summary_values),
        start=1
    ):

        header_cell = worksheet.cell(
            row=11,
            column=column,
            value=header
        )

        value_cell = worksheet.cell(
            row=12,
            column=column,
            value=value
        )


        header_cell.font = Font(
            bold=True,
            size=10,
            color=WHITE
        )

        header_cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        value_cell.font = Font(
            bold=True,
            size=15
        )

        value_cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )


        header_cell.border = thin_border
        value_cell.border = thin_border


        if column == 1:

            header_cell.fill = PatternFill(
                'solid',
                fgColor=GREEN
            )

            value_cell.fill = PatternFill(
                'solid',
                fgColor=GREEN_LIGHT
            )

            value_cell.font = Font(
                bold=True,
                size=15,
                color=GREEN
            )


        elif column == 2:

            header_cell.fill = PatternFill(
                'solid',
                fgColor=RED
            )

            value_cell.fill = PatternFill(
                'solid',
                fgColor=RED_LIGHT
            )

            value_cell.font = Font(
                bold=True,
                size=15,
                color=RED
            )


        else:

            header_cell.fill = PatternFill(
                'solid',
                fgColor=INDIGO
            )

            value_cell.fill = PatternFill(
                'solid',
                fgColor=INDIGO_LIGHT
            )

            value_cell.font = Font(
                bold=True,
                size=15,
                color=INDIGO
            )


        value_cell.number_format = (
            '₹#,##0.00'
        )


    worksheet.row_dimensions[11].height = 22
    worksheet.row_dimensions[12].height = 30


    # =========================================================
    # CATEGORY SUMMARY
    # =========================================================

    worksheet['A14'] = 'Category-wise Summary'

    worksheet['A14'].font = Font(
        bold=True,
        size=13,
        color=INDIGO_DARK
    )


    category_headers = [
        'Category',
        'Type',
        'Total'
    ]


    for column, header in enumerate(
        category_headers,
        start=1
    ):

        cell = worksheet.cell(
            row=15,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True,
            color=WHITE
        )

        cell.fill = PatternFill(
            'solid',
            fgColor=INDIGO
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        cell.border = thin_border


    row = 16


    for item in category_summary:

        category_cell = worksheet.cell(
            row=row,
            column=1,
            value=item['category__name']
        )


        type_cell = worksheet.cell(
            row=row,
            column=2,
            value=item['category__type'].title()
        )


        amount_cell = worksheet.cell(
            row=row,
            column=3,
            value=float(item['total'])
        )


        category_cell.border = thin_border
        type_cell.border = thin_border
        amount_cell.border = thin_border


        category_cell.alignment = Alignment(
            horizontal='left',
            vertical='center'
        )


        type_cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )


        amount_cell.alignment = Alignment(
            horizontal='right',
            vertical='center'
        )


        amount_cell.number_format = (
            '₹#,##0.00'
        )


        # Alternating row background

        if row % 2 == 0:

            for column in range(1, 4):

                worksheet.cell(
                    row=row,
                    column=column
                ).fill = PatternFill(
                    'solid',
                    fgColor=ROW_LIGHT
                )


        # Type colors

        if item['category__type'] == 'income':

            type_cell.font = Font(
                bold=True,
                color=GREEN
            )

            amount_cell.font = Font(
                bold=True,
                color=GREEN
            )

        else:

            type_cell.font = Font(
                bold=True,
                color=RED
            )

            amount_cell.font = Font(
                bold=True,
                color=RED
            )


        row += 1


    # =========================================================
    # EMPTY REPORT STATE
    # =========================================================

    if row == 16:

        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=3
        )


        empty_cell = worksheet.cell(
            row=row,
            column=1
        )

        empty_cell.value = (
            'No transactions found for this period.'
        )

        empty_cell.font = Font(
            italic=True,
            color=TEXT_MUTED
        )

        empty_cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

        empty_cell.fill = PatternFill(
            'solid',
            fgColor=INDIGO_LIGHT
        )

        empty_cell.border = thin_border


    # =========================================================
    # TABLE FILTER
    # =========================================================

    if row > 16:

        worksheet.auto_filter.ref = (
            f'A15:C{row - 1}'
        )


    # =========================================================
    # FOOTER
    # =========================================================

    footer_row = row + 2


    worksheet.merge_cells(
        start_row=footer_row,
        start_column=1,
        end_row=footer_row,
        end_column=3
    )


    footer_cell = worksheet.cell(
        row=footer_row,
        column=1
    )


    footer_cell.value = (
        'Expense Tracker • '
        'Generated financial report'
    )


    footer_cell.font = Font(
        italic=True,
        size=9,
        color=TEXT_MUTED
    )


    footer_cell.alignment = Alignment(
        horizontal='center'
    )


    # =========================================================
    # COLUMN WIDTHS
    # =========================================================

    worksheet.column_dimensions['A'].width = 30
    worksheet.column_dimensions['B'].width = 22
    worksheet.column_dimensions['C'].width = 20


    # =========================================================
    # PRINT AREA
    # =========================================================

    worksheet.print_area = (
        f'A1:C{footer_row}'
    )


    # =========================================================
    # RESPONSE
    # =========================================================

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument'
            '.spreadsheetml.sheet'
        )
    )


    filename = (
        f'Expense_Tracker_'
        f'{report_type.title()}_Report.xlsx'
    )


    response['Content-Disposition'] = (
        f'attachment; filename="{filename}"'
    )


    workbook.save(response)


    # =========================================================
    # SAVE REPORT HISTORY
    # =========================================================

    Report.objects.create(
        user=request.user,
        report_type=report_type,
        export_format='excel',
        start_date=start_date,
        end_date=end_date
    )


    return response

@login_required
def verify_delete_account_otp(request):
    user_id = request.session.get(
        'delete_account_user_id'
    )

    if not user_id or user_id != request.user.id:
        messages.error(
            request,
            'Your account deletion verification session has expired.'
        )
        return redirect('profile')

    user = get_object_or_404(
        User,
        pk=user_id,
        is_active=True,
    )

    if request.method == 'POST':
        otp = request.POST.get(
            'otp',
            ''
        ).strip()

        if not otp.isdigit() or len(otp) != 6:
            messages.error(
                request,
                'Please enter a valid 6-digit OTP.'
            )

            return render(
                request,
                'tracker/verify_otp.html',
                {
                    'purpose': 'account_delete',
                    'email': user.email,
                }
            )

        success, message, otp_record = verify_otp(
            email=user.email,
            purpose='account_delete',
            otp=otp,
        )

        if success:
            with db_transaction.atomic():
                user.delete()

            request.session.pop(
                'delete_account_user_id',
                None
            )

            logout(request)

            messages.success(
                request,
                'Your account and associated data have been permanently deleted.'
            )

            return redirect('login')

        messages.error(
            request,
            message
        )

    return render(
        request,
        'tracker/verify_otp.html',
        {
            'purpose': 'account_delete',
            'email': user.email,
        }
    )